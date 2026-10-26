"""Quantamental Investing and AI - Excel model: universe, weights, master results, growth series."""
import json
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

NAVY = "1F3864"
GOLD = "BF8F00"
HDR_FILL = PatternFill("solid", fgColor=NAVY)
HDR_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color=NAVY)

wb = Workbook()

def sheet(title, df, note=None):
    ws = wb.create_sheet(title)
    ws.append([title])
    ws["A1"].font = TITLE_FONT
    if note:
        ws.append([note])
        ws.append([])
    hdr_row = ws.max_row + 1
    ws.append(list(df.columns))
    for c in ws[hdr_row]:
        c.fill = HDR_FILL
        c.font = HDR_FONT
    for _, row in df.iterrows():
        ws.append(list(row))
    for col in range(1, len(df.columns) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.column_dimensions["A"].width = 26
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)
    return ws

# Universe
import sys; sys.path.insert(0, "scripts")
from engine import SHORT, NAMES, SECTORS
import sys
sys.path.insert(0, "scripts")
uni = pd.DataFrame({"Ticker": [t + ".TO" for t in SHORT],
                    "Company": [NAMES[t] for t in SHORT],
                    "Sector": [SECTORS[t] for t in SHORT]})
wb.remove(wb.active)
sheet("Universe", uni, "10 TSX stocks, weekly prices in CAD, Jun 2016 - Oct 2026.")

# Static weights
w = pd.read_csv("data/static_weights.csv", index_col=0)
w.index.name = "Strategy"
sheet("Static weights", w.reset_index(),
      "Estimated on train window Jun 2016 - Dec 2021. Caps: 20% per name, 40% per sector.")

# Master results
with open("data/static_results.json") as f:
    sr = json.load(f)["results"]
with open("data/dynamic_results.json") as f:
    dr = json.load(f)
with open("data/lstm_results.json") as f:
    lr = json.load(f)
dr = {**dr, "LSTM": lr["metrics"]}
rows = []
for name, m in {**sr, **dr}.items():
    rows.append({"Strategy": name,
                 "Ann return": m.get("ann_return"),
                 "Ann vol": m.get("ann_vol"),
                 "Sharpe": m.get("sharpe"),
                 "Max drawdown": m.get("max_drawdown"),
                 "CVaR95 weekly": m.get("cvar95_weekly"),
                 "Sortino": m.get("sortino"),
                 "Calmar": m.get("calmar"),
                 "Annual turnover": m.get("annual_turnover"),
                 "Total cost drag": m.get("total_cost_drag")})
res = pd.DataFrame(rows)
sheet("Results test window", res,
      "Test window Jan 2022 - Oct 2026, net of 10 bps trading costs. Risk-free rate 2%.")

# Growth
sg = pd.read_csv("data/static_growth.csv", index_col=0, parse_dates=True)
dg = pd.read_csv("data/dynamic_growth.csv", index_col=0, parse_dates=True)
lg = pd.read_csv("data/lstm_growth.csv", index_col=0, parse_dates=True)
growth = pd.concat([sg, dg.drop(columns=["S&P/TSX Benchmark"], errors="ignore"), lg], axis=1)
growth.index.name = "Week"
sheet("Growth of 100", growth.reset_index(),
      "Growth of $100, test window, net of costs.")

# Betas
with open("data/static_results.json") as f:
    betas = json.load(f)["betas"]
b = pd.DataFrame({"Stock": list(betas.keys()), "CAPM beta (train)": list(betas.values())})
sheet("CAPM betas", b, "OLS betas vs S&P/TSX, train window, excess returns both sides.")

wb.save("Quantamental_Investing_and_AI_Portfolio_Model.xlsx")
print("saved Quantamental_Investing_and_AI_Portfolio_Model.xlsx")
